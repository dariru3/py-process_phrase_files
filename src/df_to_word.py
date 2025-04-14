from docx import Document
import os
import pandas as pd
from datetime import datetime
from .process_mxliff import parse_mxliff_to_df
from .format_helper import apply_formatting_pipe
from .merge_df import merge_dfs
from .process_word import process_word_file
from .save_formatting import reapply_formatting_to_column
from .config_loader import CONFIG

def delete_column_in_table(table, column_index):
    grid = table._tbl.find("w:tblGrid", table._tbl.nsmap)
    for cell in table.column_cells(column_index):
        cell._tc.getparent().remove(cell._tc)
    col_elem = grid[column_index]
    grid.remove(col_elem)

def get_file_pairs(folder_path):
    docx_files = {}
    mxliff_files = {}
    for filename in os.listdir(folder_path):
        base_name, ext = os.path.splitext(filename)
        if ext == ".docx":
            docx_files[base_name] = filename
        elif ext == ".mxliff":
            mxliff_files[base_name] = filename

    # Pairing files with the same base name
    pairs = []
    for base_name, docx_file in docx_files.items():
        mxliff_file = mxliff_files.get(base_name)
        if mxliff_file:
            pairs.append((docx_file, mxliff_file))
    return pairs

def dataframe_to_word_table(docx_file, df, output_folder, formatting_info):
    t_settings = CONFIG["TableFormattingSettings"]
    # Create new .docx file with a new blank table
    doc = Document()
    table = doc.add_table(rows=1, cols=len(df.columns))
    table.autofit = False

    # Rename column headers
    df.rename(columns=t_settings["NewColumnNames"], inplace=True)
    # Reassign the 'p' column from id numbers to index numbers
    df['p'] = range(1, len(df) + 1)

    # Add header row to new table
    for i, column in enumerate(df.columns):
        table.cell(0, i).text = str(column)

    # Add data rows to new table
    for i, row in df.iterrows():
        cells = table.add_row().cells
        for j, value in enumerate(row):
            if pd.isnull(value) or value == "None":
                cells[j].text = ""
            else:
                cells[j].text = str(value)

    apply_formatting_pipe(table, doc)

    # Reapply formatting to Enlglish text
    reapply_formatting_to_column(table=table, table_num=0, col_num=2, formatting_info=formatting_info)

    # Drop the 'Match' column after all formatting is done
    match_column_index = CONFIG["ConditionalFormattingSettings"]["MatchColumnIndex"]
    delete_column_in_table(table, match_column_index)

    # Save new .docx file
    # Check output folder
    if not os.path.exists(output_folder):
        os.makedirs(output_folder)

    # Use original .docx file name, add "_merged"
    output_file_path = os.path.join(output_folder, f"{os.path.splitext(docx_file)[0]}_merged.docx")

    # Add metadata
    doc.core_properties.created = datetime.now()
    doc.core_properties.modified = datetime.now()

    # Save file and notify in console
    doc.save(output_file_path)
    print(f"Merged tables saved as Word document: {output_file_path}.")

def process_files(docx_file, mxliff_file, input_folder, output_folder):
    df_word, formatting_info = None, None

    # Process the Word and MXLIFF files
    df_word, formatting_info = process_word_file(os.path.join(input_folder, docx_file), output_folder)

    # Check Word data, then process MXLIFF files
    if df_word is not None and not df_word.empty:
        df_mxliff = parse_mxliff_to_df(os.path.join(input_folder, mxliff_file))
    else:
        print("Failed to process Word file.")
        return

    # Merge the DataFrames
    merged_df = merge_dfs(df_word, df_mxliff)

    # Save the merged DataFrame to a Word document
    dataframe_to_word_table(docx_file, merged_df, output_folder, formatting_info)
